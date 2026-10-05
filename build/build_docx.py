#!/usr/bin/env python3
"""Build Word versions of all Markdown pages in docs/.

    python3 build/build_docx.py                    # all pages -> docs/downloads/
    python3 build/build_docx.py docs/essen.md      # single page
    SITE_URL=https://example.org/onboarding/ python3 build/build_docx.py

Pipeline per page:
  1. pandoc converts Markdown to DOCX with templates/reference.docx (styles)
     and build/filters/docx.lua (form fields, links, line breaks).
  2. This script post-processes the DOCX with python-docx: header, footer with
     page numbers, table widths, header-row shading and form layout.

Pages with `docx: false` in their front matter are skipped.
"""

import argparse
import datetime as dt
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.run import Run

from ooxml_fix import fix_docx

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
REFERENCE = ROOT / "templates" / "reference.docx"
FILTER = ROOT / "build" / "filters" / "docx.lua"

TEXT_WIDTH_TW = 9638          # 17 cm in twips (A4 minus 2 x 2 cm)
HEADER_FILL = "D6E3F3"        # 20 % NCT blue
LABEL_FILL = "F4F6F9"
BORDER = "BFC7D5"

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def front_matter(path: Path) -> dict:
    text = path.read_text(encoding="utf-8")
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
    if not m:
        return {}
    return yaml.safe_load(m.group(1)) or {}


def version_date(path: Path) -> str:
    try:
        out = subprocess.run(
            ["git", "log", "-1", "--format=%cs", "--", str(path)],
            cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()
        if out:
            return out
    except (subprocess.CalledProcessError, FileNotFoundError):
        pass
    return dt.date.today().isoformat()


def add_namespaces(docx_path: Path):
    """Declare w14/mc in document.xml so that checkbox content controls parse."""
    tmp = docx_path.with_suffix(".tmp")
    with zipfile.ZipFile(docx_path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                root_end = xml.index(">", xml.index("<w:document"))
                root = xml[:root_end]
                extra = ""
                if "xmlns:w14=" not in root:
                    extra += f' xmlns:w14="{W14}"'
                if "xmlns:mc=" not in root:
                    extra += f' xmlns:mc="{MC}"'
                if "mc:Ignorable=" not in root:
                    extra += ' mc:Ignorable="w14"'
                xml = root + extra + xml[root_end:]
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(docx_path)


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_cell_fill(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    tcpr.append(el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))


def field(paragraph, instr):
    fld = el("w:fldSimple", **{"w:instr": instr})
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


# --------------------------------------------------------------------------- #
# tables
# --------------------------------------------------------------------------- #
def column_widths(table, is_form):
    ncols = len(table.columns)
    if is_form and ncols == 2:
        return [0.40, 0.60]
    weights = []
    for ci in range(ncols):
        longest = 4
        for row in table.rows:
            try:
                txt = row.cells[ci].text
            except IndexError:
                continue
            longest = max(longest, min(len(txt), 70))
        weights.append(longest ** 0.75)
    total = sum(weights)
    widths = [w / total for w in weights]
    floor = 0.12 if ncols > 2 else 0.25
    widths = [max(w, floor) for w in widths]
    total = sum(widths)
    return [w / total for w in widths]


def style_table(table):
    tbl = table._tbl
    tblpr = tbl.tblPr
    is_form = next(tbl.iter(qn("w:sdt")), None) is not None

    # full width, fixed layout
    for tag in ("w:tblW", "w:tblLayout", "w:tblBorders", "w:tblCellMar", "w:jc"):
        for old in tblpr.findall(qn(tag)):
            tblpr.remove(old)
    tblpr.append(el("w:tblW", **{"w:w": TEXT_WIDTH_TW, "w:type": "dxa"}))
    tblpr.append(el("w:tblLayout", **{"w:type": "fixed"}))
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(el(f"w:{side}", **{"w:val": "single", "w:sz": 4,
                                           "w:space": 0, "w:color": BORDER}))
    tblpr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", 45), ("left", 85), ("bottom", 45), ("right", 85)):
        mar.append(el(f"w:{side}", **{"w:w": val, "w:type": "dxa"}))
    tblpr.append(mar)

    widths = [round(w * TEXT_WIDTH_TW) for w in column_widths(table, is_form)]
    grid = tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        grid.append(el("w:gridCol", **{"w:w": w}))

    for ri, row in enumerate(table.rows):
        header = ri == 0
        if header:
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn("w:tblHeader")) is None:
                trpr.append(el("w:tblHeader"))
        elif is_form:
            row.height = Cm(0.85)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            if "Foto" in row._tr.xml and "PlaceholderText" in row._tr.xml:
                row.height = Cm(4.5)
        for ci, cell in enumerate(row.cells):
            cell.width = widths[ci] if ci < len(widths) else widths[-1]
            tcw = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tcw.set(qn("w:type"), "dxa")
            tcw.set(qn("w:w"), str(widths[min(ci, len(widths) - 1)]))
            if header:
                set_cell_fill(cell, HEADER_FILL)
            elif is_form and ci == 0:
                set_cell_fill(cell, LABEL_FILL)
            cell.vertical_alignment = (WD_CELL_VERTICAL_ALIGNMENT.CENTER if is_form
                                       else WD_CELL_VERTICAL_ALIGNMENT.TOP)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                # all runs, including those inside hyperlinks and content controls
                for r_el in p._p.iter(qn("w:r")):
                    r = Run(r_el, p)
                    r.font.size = Pt(9)
                    if header:
                        r.font.bold = True


def spacing_after_blocks(doc):
    """Give the first paragraph after a table or a list some space before."""
    body = doc.element.body
    prev = None
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            is_list = child.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None
            prev_is_list = (prev is not None and prev.tag == qn("w:p") and
                            prev.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None)
            prev_is_table = prev is not None and prev.tag == qn("w:tbl")
            if (prev_is_table or prev_is_list) and not is_list:
                ppr = child.find(qn("w:pPr"))
                if ppr is None:
                    ppr = OxmlElement("w:pPr")
                    child.insert(0, ppr)
                if ppr.find(qn("w:spacing")) is None:
                    ppr.append(el("w:spacing", **{"w:before": 120}))
        prev = child


# --------------------------------------------------------------------------- #
# header / footer
# --------------------------------------------------------------------------- #
def header_footer(doc, title, version, site_url):
    sec = doc.sections[0]
    right = Cm(17.0)

    sec.header.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]
    hp.style = doc.styles["Header"]
    hp.paragraph_format.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
    hp.add_run(title)
    hp.add_run("\tVersion " + version)

    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.style = doc.styles["Footer"]
    fp.paragraph_format.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
    fp.add_run("This document changes. Use the current version " + ("at " + site_url if site_url else "from the onboarding repository") + ".")
    fp.add_run("\tPage ")
    field(fp, "PAGE")
    fp.add_run(" of ")
    field(fp, "NUMPAGES")


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def build(src: Path, out_dir: Path, site_url: str | None):
    meta = front_matter(src)
    if meta.get("docx", True) is False:
        return None
    rel = src.relative_to(DOCS).as_posix()
    title = str(meta.get("title") or src.stem)
    version = version_date(src)
    dest = out_dir / Path(rel).with_suffix(".docx")
    dest.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / "raw.docx"
        cmd = ["pandoc", str(src), "-f", "markdown", "-t", "docx",
               "--reference-doc", str(REFERENCE), "--lua-filter", str(FILTER),
               "-M", f"docpath={rel}", "-o", str(raw)]
        if site_url:
            cmd += ["-M", f"site_url={site_url}"]
        subprocess.run(cmd, check=True)
        add_namespaces(raw)

        doc = Document(raw)
        for table in doc.tables:
            style_table(table)
        spacing_after_blocks(doc)
        header_footer(doc, title, version, site_url)
        cp = doc.core_properties
        cp.title = title
        cp.author = "Bressem Lab"
        cp.comments = f"Generated from {rel}. Edit the Markdown source, not this file."
        doc.save(dest)
    fix_docx(dest)
    return dest


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="*", type=Path)
    ap.add_argument("--out", type=Path, default=DOCS / "downloads")
    ap.add_argument("--site-url", default=os.environ.get("SITE_URL") or None)
    args = ap.parse_args()

    if shutil.which("pandoc") is None:
        sys.exit("pandoc not found. Install it from https://pandoc.org/installing.html")
    if not REFERENCE.exists():
        sys.exit("templates/reference.docx missing. Run build/make_reference_docx.py")

    out_dir = args.out.resolve()
    files = [f.resolve() for f in args.files] or sorted(
        p for p in DOCS.rglob("*.md") if out_dir not in p.parents)
    built = 0
    for src in files:
        dest = build(src, out_dir, args.site_url)
        if dest:
            built += 1
            print(f"  {src.relative_to(ROOT)} -> {dest.relative_to(ROOT) if ROOT in dest.parents else dest}")
    print(f"{built} Word file(s) written to {out_dir}")


if __name__ == "__main__":
    main()
